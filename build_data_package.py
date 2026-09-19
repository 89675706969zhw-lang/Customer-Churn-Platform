# -*- coding: utf-8 -*-
"""
S1 数据封装脚本
=====================================
生成 Rostelecom 客户流失预测项目的合成数据集（30 000 客户 × 24 个月），
并输出带完整来源标注的 Excel 数据手册 + CSV 原始文件。

用法 / Usage:
    python build_data_package.py

产物 / Outputs:
    data/Rostelecom_流失预测_数据手册.xlsx   —— 说明 / 来源 / 字段字典 / 校验 / 主数据
    data/raw/customer_base.csv
    data/raw/monthly_usage.csv
    data/raw/churn_events.csv
    data/raw/model_table.csv

⚠ 所有数据为合成数据（синтетические данные），锚定于 ПАО «Ростелеком» 官网公开披露指标。
"""

import os
import sys
import numpy as np
import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import generate_synthetic_data as gs

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ------------------------------------------------------------------ 参数
N = 30_000
MONTHS = 24
SEED = 42
BASE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(BASE, "data")
RAW = os.path.join(DATA, "raw")
XLSX = os.path.join(DATA, "Rostelecom_流失预测_数据手册.xlsx")

gs.OBS_MONTHS = 18
gs.LABEL_MONTHS = MONTHS - gs.OBS_MONTHS

RT = "4A00B8"
RT_LIGHT = "F6F1FF"

# ------------------------------------------------------------------ 数据来源
SOURCES = [
    # 参数, 取值, 来源类型, 来源说明, 链接
    ("移动业务用户基数", "48.9 百万", "官网披露",
     "Ростелеком 2025 年 Q4 及全年业绩（МСФО，2026-02-26 发布）",
     "https://www.company.rt.ru/ir/news_calendar/"),
    ("移动业务流失率", "7.1%（2025 Q1）", "官网披露",
     "2025 Q1 业绩「Показатели мобильного бизнеса」表；2024 Q1 为 6.4%，同比 +0.7 п.п.",
     "https://www.company.rt.ru/press/news_ir/news/d473194/"),
    ("移动业务年收入", "2 883.45 亿 ₽（+9%）", "官网披露",
     "2025 年全年业绩分业务收入", "https://www.company.rt.ru/ir/news_calendar/"),
    ("移动 ARPU（推算）", "≈491 ₽/月", "推算",
     "2 883.45 亿 ₽ ÷ 48.9 百万 ÷ 12 月；含 MVNO 与 B2B，为混合口径",
     "https://www.company.rt.ru/ir/news_calendar/"),
    ("光纤宽带家庭用户", "12.8 百万（+8%）", "官网披露",
     "2025 Q4 业绩；B2C 口径 12.0 百万（Q1，+7%）",
     "https://www.company.rt.ru/press/news_ir/news/d473194/"),
    ("宽带 B2C ARPU", "418 ₽/月（+3%）", "官网披露",
     "2025 Q1 业绩 ARPU 表", "https://www.company.rt.ru/press/news_ir/news/d473194/"),
    ("IPTV 用户", "7.6 百万（+6%）", "官网披露",
     "2025 Q1 业绩", "https://www.company.rt.ru/press/news_ir/news/d473194/"),
    ("IPTV ARPU", "334 ₽/月（+2%）", "官网披露",
     "2025 Q1 业绩 ARPU 表", "https://www.company.rt.ru/press/news_ir/news/d473194/"),
    ("固话用户", "8.1 百万（−10%）", "官网披露",
     "2025 Q4 业绩；结构性衰退，用于反推该线流失率",
     "https://www.company.rt.ru/ir/news_calendar/"),
    ("固话 ARPU", "236 ₽/月", "官网披露",
     "2025 Q1 业绩 ARPU 表", "https://www.company.rt.ru/press/news_ir/news/d473194/"),
    ("虚拟 PBX 用户", "1.4 百万（+14%）", "官网披露",
     "2025 Q4 业绩", "https://www.company.rt.ru/ir/news_calendar/"),
    ("虚拟 PBX ARPU", "774 ₽/月", "官网披露",
     "2025 Q1 业绩 ARPU 表", "https://www.company.rt.ru/press/news_ir/news/d473194/"),
    ("OIBDA 利润率", "37.9%（2024 为 38.8%）", "官网披露",
     "2025 年全年业绩；用于把收入影响折算为利润",
     "https://www.company.rt.ru/ir/news_calendar/"),
    ("全年营业收入", "8 728 亿 ₽（+12%）", "官网披露",
     "2025 年全年业绩（МСФО）",
     "https://www.akm.ru/eng/news/rostelecom-s-net-profit-under-ifrs-for-2025-fell-by-22/"),
    ("全年净利润", "187 亿 ₽（−22%）", "官网披露",
     "2025 年全年业绩；增收不增利，构成本研究的业务背景",
     "https://www.akm.ru/eng/news/rostelecom-s-net-profit-under-ifrs-for-2025-fell-by-22/"),
    ("IR 新闻日历（检索入口）", "—", "官网入口",
     "若具体链接失效，按披露日期在此检索原始文件",
     "https://www.company.rt.ru/ir/news_calendar/"),
    ("宽带线流失率", "9.0%", "假设",
     "公司未单独披露；按行业区间设定，须做敏感性分析", "—"),
    ("IPTV 线流失率", "11.0%", "假设",
     "公司未单独披露；按行业区间设定，须做敏感性分析", "—"),
    ("固话线流失率", "18.0%", "推算",
     "由用户数同比 −10% 的结构性衰退反推（净减少 = 流失 − 新增）", "—"),
    ("合约类型分布", "月付 55% / 年付 30% / 两年 15%", "假设",
     "未披露；按典型电信业务结构设定，须做敏感性分析", "—"),
    ("付费方式分布", "自动 45% / 卡 35% / 账单 20%", "假设",
     "未披露；按典型电信业务结构设定", "—"),
    ("干预成本（电话）", "800 ₽/人，成功率 25%", "假设",
     "S0 决策四确定的情景参数，用于 Uplift 排序与 ROI 仿真", "—"),
    ("干预成本（短信）", "80 ₽/人，成功率 8%", "假设", "同上", "—"),
    ("干预成本（App 优惠）", "350 ₽/人，成功率 18%", "假设", "同上", "—"),
    ("干预成本（上门）", "5 000 ₽/人，成功率 45%", "假设", "仅限 B2B 与高 CLV 客户", "—"),
]

# ------------------------------------------------------------------ 字段字典
DICT_BASE = [
    ("customer_id", "字符串", "客户唯一编号（RT-000001 …），主键", "—"),
    ("business_line", "类别", "主业务线：mobile / broadband / iptv / fixed_voice", "决定流失率基准与 ARPU 分布"),
    ("region", "类别", "所属区域（10 个）", "用于区域分组与竞争强度建模"),
    ("competition_index", "0–1 连续", "区域竞争强度，Beta(2,2) 抽样", "推高流失倾向"),
    ("tenure_months", "整数（月）", "在网时长，Gamma(2.2, 14.5) 截顶至 1–120", "在网越久越稳定（保护因素）"),
    ("contract_type", "类别", "合约类型：monthly 55% / annual 30% / two_year 15%", "月付客户流失率显著更高"),
    ("payment_method", "类别", "付费方式：auto 45% / card 35% / invoice 20%", "自动扣款为保护因素"),
    ("tech_type", "类别", "接入技术：fiber / xdsl / copper（仅宽带与固话线）", "xDSL 迁移意愿强，推高流失"),
    ("n_products", "整数 1–4", "持有的产品数量", "交叉持有越多越难离开（强保护因素）"),
    ("has_mobile", "布尔", "是否持有移动业务", "产品持有矩阵"),
    ("has_broadband", "布尔", "是否持有宽带业务", "产品持有矩阵"),
    ("has_iptv", "布尔", "是否持有 IPTV", "产品持有矩阵"),
    ("has_pbx", "布尔", "是否持有虚拟 PBX（B2B）", "高价值黏性产品"),
    ("days_to_contract_end", "整数（天）", "距合约到期天数", "到期前 60 天内为流失高发窗口"),
]
DICT_MONTH = [
    ("customer_id", "字符串", "外键，关联客户主表", "—"),
    ("month", "整数 1–24", "月份序号（M1 = 第 1 个月）", "M1–M18 为观测窗口，M19–M24 为标签窗口"),
    ("usage_gb", "浮点", "当月使用量（GB）", "含个体效应与 12 个月季节性；趋势比水平更重要"),
    ("call_minutes", "浮点", "当月通话时长（分钟），缺失 3%", "缺失值处理检验"),
    ("bill_amount", "浮点（₽）", "当月账单金额", "以各业务线 ARPU 为中心；含偶发资费上调"),
    ("support_tickets", "整数", "当月服务工单数，Poisson 抽样", "服务质量信号，推高流失"),
    ("unresolved_tickets", "整数", "当月未解决工单数", "杀伤力大于已解决工单"),
    ("outage_hours", "浮点", "当月故障时长（小时），缺失 8%", "网络质量信号"),
    ("late_payments", "整数", "当月欠费次数", "财务压力信号"),
    ("promo_active", "布尔", "当月是否处于促销期", "促销期内的留存差异"),
]
DICT_CHURN = [
    ("customer_id", "字符串", "外键", "—"),
    ("churn_flag", "0 / 1", "标签窗口内是否流失（主标签）", "由校准后的流失概率 Bernoulli 抽样生成"),
    ("churn_month", "整数 19–24", "流失发生的月份（未流失为空）", "供离散时间生存分析使用"),
    ("churn_reason", "类别", "price / quality / competitor / tech_migration / unknown",
     "按主要驱动因素归类；tech_migration 反映 xDSL→光纤的真实迁移"),
]
DICT_MODEL = [
    ("usage_mean_6m", "浮点", "观测窗口近 6 月使用量均值", "水平特征"),
    ("usage_last3m / usage_prior3m", "浮点", "近 3 月 / 前 3 月使用量均值", "用于构造下降斜率"),
    ("usage_decline_3m", "0–1", "使用量下降幅度（前 3 月 → 近 3 月）", "趋势特征，预测力强于水平特征"),
    ("tickets_6m / tickets_12m", "整数", "近 6 / 12 月工单数（截顶至 8）", "服务交互特征"),
    ("unresolved_6m", "整数", "近 6 月未解决工单数", "服务交互特征"),
    ("bill_mean_6m / bill_std_6m", "浮点", "近 6 月账单均值与波动率", "账单特征"),
    ("bill_last / bill_first", "浮点", "观测窗口末月 / 首月账单", "用于计算账单增幅"),
    ("bill_increase_pct", "−0.5 至 1.0", "窗口首尾账单增幅", "账单突增是强流失信号"),
    ("late_payments_6m", "整数", "近 6 月欠费次数（截顶至 6）", "账单特征"),
    ("outage_mean_6m", "浮点", "近 6 月平均故障时长", "网络质量特征"),
    ("promo_months", "整数", "观测窗口内处于促销的月数", "促销特征"),
]


def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor=RT)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 26
    ws.freeze_panes = f"A{row+1}"


def autosize(ws, minw=10, maxw=46):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        try:
            w = max((len(str(c.value)) for c in col if c.value is not None), default=minw)
        except Exception:
            w = minw
        ws.column_dimensions[letter].width = min(max(w + 3, minw), maxw)


def write_df(ws, df, autosize_on=True):
    ws.append(list(df.columns))
    style_header(ws, len(df.columns))
    for row in df.itertuples(index=False):
        ws.append(list(row))
    if autosize_on:
        autosize(ws)


def main():
    os.makedirs(RAW, exist_ok=True)
    rng = np.random.default_rng(SEED)
    print(f"[1/5] 生成合成数据：n={N:,}, months={MONTHS}, seed={SEED}")
    customers = gs.build_customers(N, rng)
    print("[2/5] 生成月度行为表…")
    monthly = gs.build_monthly(customers, MONTHS, rng)
    monthly = gs.inject_missing(monthly, rng)
    print("[3/5] 聚合观测窗口特征并生成标签…")
    feats = gs.aggregate_features(monthly, gs.OBS_MONTHS)
    full = gs.generate_labels(customers, feats, rng)
    hist = gs.generate_treatment_history(full, rng)

    base_cols = ["customer_id", "business_line", "region", "competition_index",
                 "tenure_months", "contract_type", "payment_method", "tech_type",
                 "n_products", "has_mobile", "has_broadband", "has_iptv", "has_pbx",
                 "days_to_contract_end"]
    label_cols = ["customer_id", "churn_flag", "churn_month", "churn_reason"]
    df_base = full[base_cols].copy()
    df_label = full[label_cols].copy()
    df_model = full[base_cols + ["churn_flag", "churn_month"]].merge(feats, on="customer_id")

    print("[4/5] 写出 CSV…")
    df_base.to_csv(os.path.join(RAW, "customer_base.csv"), index=False, encoding="utf-8-sig")
    monthly.to_csv(os.path.join(RAW, "monthly_usage.csv"), index=False, encoding="utf-8-sig")
    df_label.to_csv(os.path.join(RAW, "churn_events.csv"), index=False, encoding="utf-8-sig")
    df_model.to_csv(os.path.join(RAW, "model_table.csv"), index=False, encoding="utf-8-sig")
    hist.to_csv(os.path.join(RAW, "intervention_history.csv"), index=False, encoding="utf-8-sig")
    rate_t, rate_c = (hist.loc[hist.treated == t, "churned"].mean() for t in (1, 0))
    print(f"  干预历史: 覆盖率={hist.treated.mean():.1%}  对照={rate_c:.3f}  "
          f"干预={rate_t:.3f}  ATE≈{rate_c - rate_t:+.3f}")

    # ---- 锚点校验 ----
    rows = []
    for line, cfg in gs.LINES.items():
        sub = df_label.merge(df_base[["customer_id", "business_line"]], on="customer_id")
        sub = sub[sub["business_line"] == line]
        ids = set(sub["customer_id"])
        m = monthly[monthly["customer_id"].isin(ids)]
        rows.append({
            "业务线": line,
            "目标流失率": cfg["churn"],
            "实际流失率": round(float(sub["churn_flag"].mean()), 4),
            "偏差": round(float(sub["churn_flag"].mean() - cfg["churn"]), 4),
            "目标ARPU(₽/月)": cfg["arpu"],
            "实际ARPU(₽/月)": round(float(m["bill_amount"].mean()), 1),
            "ARPU偏差%": round(float((m["bill_amount"].mean() - cfg["arpu"]) / cfg["arpu"] * 100), 2),
            "用户数": int(len(sub)),
            "占比": round(len(sub) / N, 4),
            "校验": "OK" if abs(sub["churn_flag"].mean() - cfg["churn"]) < 0.012 else "CHECK",
        })
    df_check = pd.DataFrame(rows)

    print("[5/5] 写出 Excel 数据手册…")
    wb = Workbook()

    # --- 00 说明 ---
    ws = wb.active
    ws.title = "00_说明"
    info = [
        ["Rostelecom 客户流失预测项目 · 数据手册", ""],
        ["", ""],
        ["生成日期", "2026-09-18"],
        ["数据性质", "合成数据（синтетические данные），锚定于公司公开披露指标"],
        ["客户数量", f"{N:,}"],
        ["时间跨度", f"{MONTHS} 个月（观测窗口 M1–M{gs.OBS_MONTHS}，标签窗口 M{gs.OBS_MONTHS+1}–M{MONTHS}）"],
        ["随机种子", f"{SEED}（固定，结果完全可复现）"],
        ["生成脚本", "generate_synthetic_data.py / build_data_package.py"],
        ["", ""],
        ["工作表导航", ""],
        ["01_数据来源", "全部锚点参数的取值、来源类型与来源链接（可点击）"],
        ["02_字段字典", "customer_base / monthly_usage / churn_events / model_table 字段说明"],
        ["03_锚点校验", "生成数据的实际统计量与财报目标的对比"],
        ["04_客户主表", f"{N:,} 行客户静态属性"],
        ["05_流失标签", f"{N:,} 行流失标签与原因"],
        ["06_月度行为样例", "前 500 名客户 × 24 个月（完整 720 000 行见 CSV）"],
        ["", ""],
        ["⚠ 重要声明", ""],
        ["数据性质说明",
         "本数据集为合成数据，其统计量锚定于 ПАО «Ростелеком» 官网公开披露数据，"
         "但个体层面不包含任何真实客户信息。"],
        ["合规说明",
         "不使用任何企业内部数据或个人信息；所引用财务数据均为官网已公开披露文件，属于公开信息。"],
        ["结论边界",
         "本数据集可用于验证方法可行性与系统可用性，不能用于推断公司真实的客户流失驱动因素。"
         "凡涉及真实业务的推断须表述为「在该合成数据集上观察到……」。"],
        ["数据本地化",
         "合成数据不受 152-ФЗ《个人数据法》约束；若未来接入真实数据，数据库须部署于俄罗斯境内。"],
    ]
    for r in info:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=15, color=RT)
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, max_col=2):
        row[0].font = Font(bold=True, size=11)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 96

    # --- 01 数据来源 ---
    ws = wb.create_sheet("01_数据来源")
    ws.append(["参数", "取值", "来源类型", "来源说明", "来源链接"])
    style_header(ws, 5)
    for p, v, t, d, u in SOURCES:
        ws.append([p, v, t, d, u])
        cell = ws.cell(row=ws.max_row, column=5)
        if u.startswith("http"):
            cell.hyperlink = u
            cell.font = Font(color="0563C1", underline="single")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=4).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=1).font = Font(bold=True)
    for col, w in zip("ABCDE", [24, 26, 12, 52, 46]):
        ws.column_dimensions[col].width = w

    # --- 02 字段字典 ---
    ws = wb.create_sheet("02_字段字典")
    row = 1
    for title, dic in [("表 A · customer_base（客户主表）", DICT_BASE),
                       ("表 B · monthly_usage（月度行为表）", DICT_MONTH),
                       ("表 C · churn_events（流失事件表）", DICT_CHURN),
                       ("表 D · model_table 中的聚合特征部分", DICT_MODEL)]:
        ws.cell(row=row, column=1, value=title).font = Font(bold=True, size=12, color=RT)
        row += 1
        ws.cell(row=row, column=1, value="字段名")
        ws.cell(row=row, column=2, value="类型 / 取值")
        ws.cell(row=row, column=3, value="说明")
        ws.cell(row=row, column=4, value="建模意义")
        for c in range(1, 5):
            cell = ws.cell(row=row, column=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="7A6FA0")
        row += 1
        for f, t, d, m in dic:
            ws.cell(row=row, column=1, value=f)
            ws.cell(row=row, column=2, value=t)
            ws.cell(row=row, column=3, value=d)
            ws.cell(row=row, column=4, value=m)
            row += 1
        row += 1
    for col, w in zip("ABCD", [28, 22, 56, 44]):
        ws.column_dimensions[col].width = w

    # --- 03 锚点校验 ---
    ws = wb.create_sheet("03_锚点校验")
    write_df(ws, df_check)
    note_row = ws.max_row + 2
    ws.cell(row=note_row, column=1, value="校验判据：实际流失率与目标偏差 < 0.012 记为 OK；ARPU 偏差 < 2% 视为命中。").font = Font(italic=True, color="8A8698")
    ws.cell(row=note_row + 1, column=1,
            value="模型性能基准（n=3 000 预实验）：ROC-AUC 0.829 / PR-AUC 0.361 / Lift@Top10% 5.00 / Recall@Top10% 50.0%").font = Font(italic=True, color="8A8698")

    # --- 04/05/06 数据 ---
    ws = wb.create_sheet("04_客户主表")
    write_df(ws, df_base, autosize_on=False)
    for col, w in zip(range(1, len(df_base.columns) + 1),
                      [14, 14, 12, 16, 14, 14, 14, 12, 12, 12, 14, 12, 11, 18]):
        ws.column_dimensions[get_column_letter(col)].width = w

    ws = wb.create_sheet("05_流失标签")
    write_df(ws, df_label, autosize_on=False)
    for col, w in zip(range(1, len(df_label.columns) + 1), [14, 12, 14, 20]):
        ws.column_dimensions[get_column_letter(col)].width = w

    sample_ids = df_base["customer_id"].head(500)
    df_month_sample = monthly[monthly["customer_id"].isin(set(sample_ids))].copy()
    ws = wb.create_sheet("06_月度行为样例")
    write_df(ws, df_month_sample, autosize_on=False)
    for col, w in zip(range(1, len(df_month_sample.columns) + 1),
                      [14, 8, 12, 14, 13, 15, 17, 13, 14, 13]):
        ws.column_dimensions[get_column_letter(col)].width = w

    wb.save(XLSX)

    print("\n" + "=" * 68)
    print("锚点校验 / Anchor validation")
    print("=" * 68)
    print(df_check.to_string(index=False))
    print("=" * 68)
    print("\n产物 / Outputs:")
    print(f"  {XLSX}")
    for f in ["customer_base.csv", "monthly_usage.csv", "churn_events.csv", "model_table.csv"]:
        p = os.path.join(RAW, f)
        print(f"  {p}  ({os.path.getsize(p)/1024/1024:.1f} MB)")
    print("\n下一步: python check_signal_strength.py ./data/raw/model_table.csv")


if __name__ == "__main__":
    main()
